"""Verify CL billing column: shares sum to 1 and amounts reconcile."""
from openpyxl import load_workbook
from pathlib import Path

XLSX = Path(r"c:\Users\rober\OneDrive\Hausverwaltung_Vahrn\Abrechnung_KondominiumVitroler_V2.xlsx")
wb = load_workbook(XLSX, data_only=True)
ws = wb["Verrechnung"]

# Column CL = 90? Let's find by date
from openpyxl.utils import column_index_from_string
for col in ["CL", "CN", "CQ", "G", "O"]:
    c = column_index_from_string(col)
    inv = ws.cell(40, c).value
    hz = ws.cell(47, c).value
    ww = ws.cell(48, c).value
    hz_eg = ws.cell(52, c).value
    hz_1 = ws.cell(53, c).value
    hz_2 = ws.cell(54, c).value
    ww_eg = ws.cell(59, c).value
    ww_1 = ws.cell(60, c).value
    ww_2 = ws.cell(61, c).value
    pay_eg = ws.cell(72, c).value
    pay_1 = ws.cell(73, c).value
    pay_2 = ws.cell(74, c).value
    amount = ws.cell(40, c).value
    print(f"\n=== Col {col} amount={amount} ===")
    print(f"  Anteil HZ+WW: {hz}+{ww} = {(hz or 0)+(ww or 0)}")
    print(f"  Anteil HZ EG+1+2: {hz_eg}+{hz_1}+{hz_2} = {(hz_eg or 0)+(hz_1 or 0)+(hz_2 or 0)}")
    print(f"  Anteil WW EG+1+2: {ww_eg}+{ww_1}+{ww_2} = {(ww_eg or 0)+(ww_1 or 0)+(ww_2 or 0)}")
    total_pay = (pay_eg or 0)+(pay_1 or 0)+(pay_2 or 0)
    print(f"  zu bezahlen sum={total_pay} vs invoice={amount} diff={total_pay-(amount or 0)}")

# Bottom labels
print("\n=== Bottom area ===")
for r in range(75, 86):
    row = []
    for c in range(1, ws.max_column+1):
        v = ws.cell(r, c).value
        if v is not None:
            from openpyxl.utils import get_column_letter
            row.append(f"{get_column_letter(c)}={v!r}")
    if row:
        print(f"R{r}: {row[:20]}")
